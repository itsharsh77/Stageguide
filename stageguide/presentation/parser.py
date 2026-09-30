"""Extract presentation text locally, or print it as JSON with python -m."""

import argparse
import json
from pathlib import Path
import re
import sys
from typing import Iterable, List, Optional, Sequence, Union

import pymupdf
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .models import PresentationPage


class PresentationError(ValueError):
    """A file cannot be ingested as a supported presentation."""


class UnsupportedFileError(PresentationError):
    """The supplied filename has an unsupported extension."""


# English-style numeric literals, retaining signs, decimal points and grouping.
# A percentage contributes its numeric part to numbers, too.
_NUMBER = r"[+\-−]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?|[+\-−]?\.\d+"
_MENTION = re.compile(
    rf"(?<![\w.])(?P<number>{_NUMBER})(?!\w|\.\d)"
    r"(?P<percent>\s*(?:[%％]|percent(?:age)?\b|per\s+cent\b))?",
    re.IGNORECASE,
)


def _clean_text(text: str) -> str:
    """Collapse horizontal whitespace and blank lines, preserving text lines."""
    return "\n".join(
        cleaned for line in text.splitlines()
        if (cleaned := " ".join(line.split()))
    )


def _page(number: int, title: Optional[str], body: str, filename: str) -> PresentationPage:
    text = "\n".join(part for part in (title, body) if part)
    mentions = list(_MENTION.finditer(text))
    return PresentationPage(
        page_number=number,
        title=title,
        body_text=body,
        numbers=[match.group("number") for match in mentions],
        percentages=[" ".join(match.group().split()) for match in mentions if match.group("percent")],
        source_filename=filename,
    )


def _shape_text(shapes: Iterable, title_id: Optional[int]) -> Iterable[str]:
    for shape in shapes:
        if shape.shape_id == title_id:
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shape_text(shape.shapes, title_id)
        elif shape.has_text_frame:
            yield shape.text
        elif shape.has_table:
            for row in shape.table.rows:
                yield " | ".join(cell.text for cell in row.cells if not cell.is_spanned)


def _parse_pptx(path: Path) -> List[PresentationPage]:
    presentation = Presentation(str(path))
    pages = []
    for number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = _clean_text(title_shape.text) or None if title_shape is not None else None
        title_id = title_shape.shape_id if title_shape is not None else None
        body = _clean_text("\n".join(_shape_text(slide.shapes, title_id)))
        pages.append(_page(number, title, body, path.name))
    return pages


def _parse_pdf(path: Path) -> List[PresentationPage]:
    pages = []
    with pymupdf.open(str(path)) as document:
        if not document.is_pdf:
            raise PresentationError(f"File is not a PDF: {path.name}")
        if document.needs_pass:
            raise PresentationError(f"Password-protected PDF is not supported: {path.name}")
        for number, page in enumerate(document, start=1):
            # Exclude image bytes; only text and font information are needed.
            blocks = page.get_text("dict", sort=True, flags=pymupdf.TEXTFLAGS_TEXT)["blocks"]
            lines = []
            for block in blocks:
                if block["type"] != 0:
                    continue
                for line in block["lines"]:
                    spans = line["spans"]
                    text = _clean_text("".join(span["text"] for span in spans))
                    if text:
                        size = max(span["size"] for span in spans if span["text"].strip())
                        lines.append((text, size, line["bbox"][1]))

            title = None
            if len(lines) > 1:
                first_text, first_size, first_y = lines[0]
                remaining_size = max(line[1] for line in lines[1:])
                # Only a short, distinctly larger first line near the top is a title.
                if (first_y <= page.rect.height * 0.35 and len(first_text) <= 200
                        and first_size >= remaining_size * 1.2):
                    title = first_text
            body = "\n".join(line[0] for line in (lines[1:] if title else lines))
            pages.append(_page(number, title, body, path.name))
    return pages


def parse_presentation(path: Union[str, Path]) -> List[PresentationPage]:
    """Return one record per slide/page, including empty ones.

    Raises UnsupportedFileError for extensions other than .pptx/.pdf and
    PresentationError for missing, unreadable, encrypted or corrupted input.
    No OCR, network access or embedded chart/image analysis is performed.
    """
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in {".pptx", ".pdf"}:
        raise UnsupportedFileError(f"Unsupported file type '{suffix or '(none)'}'; expected .pptx or .pdf")
    try:
        if not path.is_file():
            raise PresentationError(f"File does not exist or is not a regular file: {path}")
        return _parse_pptx(path) if suffix == ".pptx" else _parse_pdf(path)
    except PresentationError:
        raise
    except Exception as exc:
        # Third-party readers expose varied errors for damaged containers/XML.
        # Keep that backend detail as the cause while providing one public error.
        raise PresentationError(f"Cannot parse '{path.name}': file is unreadable or corrupted") from exc


def main(argv: Optional[Sequence[str]] = None) -> int:
    cli = argparse.ArgumentParser(description="Extract PPTX/PDF slides as local JSON data.")
    cli.add_argument("path", type=Path, help="Path to a .pptx or .pdf presentation")
    args = cli.parse_args(argv)
    try:
        pages = parse_presentation(args.path)
    except PresentationError as exc:
        print(json.dumps({"error": str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 1
    print(json.dumps([page.to_dict() for page in pages], indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())

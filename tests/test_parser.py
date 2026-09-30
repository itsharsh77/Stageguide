"""Exercise real presentation files and the public Python/CLI interfaces."""

import json
import subprocess
import sys
from zipfile import ZipFile

import pymupdf
from pptx import Presentation
from pptx.util import Inches
import pytest

from stageguide.presentation.models import PresentationPage
from stageguide.presentation.parser import (
    PresentationError,
    UnsupportedFileError,
    parse_presentation,
)


@pytest.fixture
def pptx_path(tmp_path):
    path = tmp_path / "test.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Growth 2026"
    slide.placeholders[1].text = "  Revenue   1,250.50\nGrowth 12.5%\nTarget 20 percent  "
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "Next steps: 3"
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(path)
    return path


@pytest.fixture
def pdf_path(tmp_path):
    path = tmp_path / "test.pdf"
    with pymupdf.open() as doc:
        page = doc.new_page()
        # Write the body first to ensure extraction uses geometric order.
        page.insert_text((50, 110), "Revenue 1,250.50\nGrowth 12.5%\nTarget 20 percent", fontsize=12)
        page.insert_text((50, 50), "Growth 2026", fontsize=24)
        page = doc.new_page()
        page.insert_text((50, 50), "Next steps: 3", fontsize=12)
        doc.new_page()
        doc.save(path)
    return path


@pytest.mark.parametrize("fixture", ["pptx_path", "pdf_path"])
def test_boundaries_and_structured_text(fixture, request):
    path = request.getfixturevalue(fixture)
    pages = parse_presentation(path)
    assert len(pages) == 3
    assert all(isinstance(page, PresentationPage) for page in pages)
    assert [page.page_number for page in pages] == [1, 2, 3]
    assert pages[0].to_dict() == {
        "page_number": 1,
        "title": "Growth 2026",
        "body_text": "Revenue 1,250.50\nGrowth 12.5%\nTarget 20 percent",
        "numbers": ["2026", "1,250.50", "12.5", "20"],
        "percentages": ["12.5%", "20 percent"],
        "source_filename": path.name,
    }
    assert pages[1].title is None
    assert pages[1].body_text == "Next steps: 3"
    assert pages[1].numbers == ["3"]
    assert pages[2].to_dict() == {
        "page_number": 3, "title": None, "body_text": "", "numbers": [],
        "percentages": [], "source_filename": path.name,
    }


def test_pptx_tables_groups_and_notes(tmp_path):
    path = tmp_path / "shapes.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "Grouped 7"
    nested = group.shapes.add_group_shape()
    nested.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "Nested 8"
    table = slide.shapes.add_table(2, 2, 0, Inches(2), Inches(4), Inches(2)).table
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "Merged 9"
    table.cell(1, 0).text = "Sales"
    table.cell(1, 1).text = "42%"
    slide.notes_slide.notes_text_frame.text = "Private notes 999"
    deck.save(path)
    page = parse_presentation(path)[0]
    assert page.body_text == "Grouped 7\nNested 8\nMerged 9\nSales | 42%"
    assert page.numbers == ["7", "8", "9", "42"]
    assert page.percentages == ["42%"]


@pytest.mark.parametrize("text,numbers,percentages", [
    ("Values -12 +3.50 −2 .75 +.5 1,000,000", ["-12", "+3.50", "−2", ".75", "+.5", "1,000,000"], []),
    ("12% 12 % 12 percent 12 per cent 12 PERCENTAGE 12％", ["12"] * 6,
     ["12%", "12 %", "12 percent", "12 per cent", "12 PERCENTAGE", "12％"]),
    ("IDs Q3 abc123 123abc and no numbers", [], []),
    ("Repeat 7, 7 and 7.", ["7", "7", "7"], []),
    ("Rate -2.5% and .5 percent", ["-2.5", ".5"], ["-2.5%", ".5 percent"]),
])
def test_numeric_mentions(tmp_path, text, numbers, percentages):
    path = tmp_path / "numbers.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, Inches(8), Inches(4)).text = text
    deck.save(path)
    page = parse_presentation(path)[0]
    assert page.numbers == numbers
    assert page.percentages == percentages


def test_empty_title_and_line_cleanup(tmp_path):
    path = tmp_path / "empty-title.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "  "
    slide.placeholders[1].text = " Café\t  team\vSecond\n\n  Third\u00a0line "
    deck.save(path)
    page = parse_presentation(path)[0]
    assert page.title is None
    assert page.body_text == "Café team\nSecond\nThird line"


@pytest.mark.parametrize("title_y,title_size,body_size", [(50, 12, 12), (400, 24, 12)])
def test_pdf_ambiguous_heading_stays_in_body(tmp_path, title_y, title_size, body_size):
    path = tmp_path / "ambiguous.pdf"
    with pymupdf.open() as doc:
        page = doc.new_page()
        page.insert_text((50, title_y), "Heading", fontsize=title_size)
        page.insert_text((50, title_y + 60), "Body 10", fontsize=body_size)
        doc.save(path)
    page = parse_presentation(path)[0]
    assert page.title is None
    assert page.body_text == "Heading\nBody 10"


def test_image_only_pdf_keeps_page(tmp_path):
    path = tmp_path / "image.pdf"
    with pymupdf.open() as doc:
        page = doc.new_page()
        pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 10, 10), False)
        pixmap.clear_with(255)
        page.insert_image(pymupdf.Rect(50, 50, 150, 150), pixmap=pixmap)
        doc.save(path)
    pages = parse_presentation(path)
    assert len(pages) == 1
    assert pages[0].body_text == ""
    assert pages[0].title is None


@pytest.mark.parametrize("fixture", ["pptx_path", "pdf_path"])
def test_uppercase_suffix_and_string_path(fixture, request):
    path = request.getfixturevalue(fixture)
    uppercase = path.with_suffix(path.suffix.upper())
    path.rename(uppercase)
    assert len(parse_presentation(str(uppercase))) == 3


@pytest.mark.parametrize("filename", ["file.txt", "file.ppt", "file.docx", "no_suffix"])
def test_unsupported_file(tmp_path, filename):
    with pytest.raises(UnsupportedFileError, match="Unsupported file type"):
        parse_presentation(tmp_path / filename)


@pytest.mark.parametrize("suffix", [".pptx", ".pdf"])
@pytest.mark.parametrize("data", [b"", b"this is not a presentation"])
def test_corrupted_files(tmp_path, suffix, data):
    path = tmp_path / ("corrupt" + suffix)
    path.write_bytes(data)
    with pytest.raises(PresentationError, match="unreadable or corrupted"):
        parse_presentation(path)


def test_invalid_pptx_container(tmp_path):
    path = tmp_path / "invalid.pptx"
    with ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "<invalid")
    with pytest.raises(PresentationError, match="unreadable or corrupted"):
        parse_presentation(path)


@pytest.mark.parametrize("directory", [False, True])
def test_missing_file_or_directory(tmp_path, directory):
    path = tmp_path / "missing.pdf"
    if directory:
        path.mkdir()
    with pytest.raises(PresentationError, match="not a regular file"):
        parse_presentation(path)


def test_password_protected_pdf(tmp_path):
    path = tmp_path / "encrypted.pdf"
    with pymupdf.open() as doc:
        doc.new_page()
        doc.save(path, encryption=pymupdf.PDF_ENCRYPT_AES_256,
                 owner_pw="owner", user_pw="secret")
    with pytest.raises(PresentationError, match="Password-protected"):
        parse_presentation(path)


def test_unreadable_file_becomes_public_error(pptx_path, monkeypatch):
    def fail(*args):
        raise PermissionError("permission denied")
    monkeypatch.setattr("stageguide.presentation.parser.Presentation", fail)
    with pytest.raises(PresentationError, match="unreadable") as error:
        parse_presentation(pptx_path)
    assert isinstance(error.value.__cause__, PermissionError)


@pytest.mark.parametrize("fixture", ["pptx_path", "pdf_path"])
def test_cli_success(fixture, request):
    path = request.getfixturevalue(fixture)
    result = subprocess.run(
        [sys.executable, "-m", "stageguide.presentation.parser", str(path)],
        capture_output=True, text=True,
    )
    assert result.returncode == 0
    assert result.stderr == ""
    assert json.loads(result.stdout) == [page.to_dict() for page in parse_presentation(path)]


@pytest.mark.parametrize("filename", ["bad.txt", "bad.pdf", "bad.pptx", "missing.pdf"])
def test_cli_error(tmp_path, filename):
    path = tmp_path / filename
    if filename != "missing.pdf":
        path.write_bytes(b"corrupted")
    result = subprocess.run(
        [sys.executable, "-m", "stageguide.presentation.parser", str(path)],
        capture_output=True, text=True,
    )
    assert result.returncode == 1
    assert result.stdout == ""
    assert json.loads(result.stderr)["error"]
    assert "Traceback" not in result.stderr


def test_to_dict_is_independent(pptx_path):
    page = parse_presentation(pptx_path)[0]
    data = page.to_dict()
    data["numbers"].append("999")
    assert "999" not in page.numbers
